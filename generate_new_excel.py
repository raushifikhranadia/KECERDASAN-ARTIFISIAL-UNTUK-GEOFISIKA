# -*- coding: utf-8 -*-
import random
import datetime
import os
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

random.seed(42)

dirs = [
    r"C:\Users\Digitalisasi\Documents\pemrograman\html_dashboard",
    r"C:\Users\Digitalisasi\Documents\pemrograman\html_dashboard\html_dashboard"
]

products = [
    ("Enterprise Hardware", "Server", "Dell", "PowerEdge R750 Rack Server", 4250.0, 0.65),
    ("Enterprise Hardware", "Server", "Lenovo", "ThinkSystem SR650 V2", 3980.0, 0.66),
    ("Enterprise Hardware", "Networking", "Cisco", "Catalyst 9300 48-Port Switch", 2850.0, 0.62),
    ("Enterprise Hardware", "Storage", "Dell", "PowerVault ME5024 SAN/NAS", 5200.0, 0.68),
    ("Enterprise Hardware", "Security", "Fortinet", "FortiGate 100F Next-Gen Firewall", 1950.0, 0.60),
    
    ("Smart Workstation", "Display", "Dell", "UltraSharp 38 Curved WQHD+", 1350.0, 0.68),
    ("Smart Workstation", "Display", "Apple", "Studio Display 27 5K Retina", 1599.0, 0.72),
    ("Smart Workstation", "Display", "Samsung", "ViewFinity S9 5K Monitor", 1299.0, 0.67),
    ("Smart Workstation", "Furniture", "Herman Miller", "Aeron Ergonomic Task Chair", 1450.0, 0.58),
    ("Smart Workstation", "Furniture", "Steelcase", "Gesture Ergonomic Executive Chair", 1380.0, 0.60),
    ("Smart Workstation", "Desk", "Jarvis", "Bamboo Smart Standing Desk 72x30", 890.0, 0.63),
    
    ("Personal Computing", "Laptop", "Apple", "MacBook Pro 16 M3 Max", 3499.0, 0.74),
    ("Personal Computing", "Laptop", "Dell", "XPS 15 OLED Touch Workstation", 2399.0, 0.70),
    ("Personal Computing", "Laptop", "Lenovo", "ThinkPad X1 Carbon Gen 12", 2150.0, 0.68),
    ("Personal Computing", "Laptop", "Asus", "ROG Zephyrus G16 Creator Edition", 2299.0, 0.71),
    ("Personal Computing", "Tablet", "Apple", "iPad Pro 13 M4 Ultra Retina", 1299.0, 0.73),
    ("Personal Computing", "Tablet", "Samsung", "Galaxy Tab S9 Ultra AMOLED", 1199.0, 0.69),
    
    ("Audio & Conferencing", "Audio", "Sony", "WH-1000XM5 ANC Headset", 399.0, 0.58),
    ("Audio & Conferencing", "Audio", "Apple", "AirPods Max Spatial Audio", 549.0, 0.65),
    ("Audio & Conferencing", "Video", "Logitech", "Rally Bar Mini Video Bar", 2499.0, 0.62),
    ("Audio & Conferencing", "Video", "Logitech", "Brio 4K Ultra HD StreamCam", 199.0, 0.55),
    ("Audio & Conferencing", "Microphone", "Shure", "MV7+ Podcast & Conf XLR/USB Mic", 279.0, 0.56),
    
    ("Accessories & Power", "Docking", "CalDigit", "TS4 Thunderbolt 4 18-Port Dock", 399.0, 0.60),
    ("Accessories & Power", "Peripherals", "Logitech", "MX Master 3S Ergonomic Mouse", 99.0, 0.52),
    ("Accessories & Power", "Peripherals", "Logitech", "MX Mechanical Wireless Keyboard", 169.0, 0.54),
    ("Accessories & Power", "Power", "APC", "Smart-UPS 1500VA LCD Battery", 650.0, 0.64),
    ("Accessories & Power", "Power", "Anker", "Prime 250W GaN Charging Station", 179.0, 0.53),
]

regions_data = [
    ("West", "California", "San Francisco"),
    ("West", "California", "Los Angeles"),
    ("West", "Washington", "Seattle"),
    ("West", "Colorado", "Denver"),
    ("Central", "Texas", "Austin"),
    ("Central", "Texas", "Dallas"),
    ("Central", "Illinois", "Chicago"),
    ("East", "New York", "New York City"),
    ("East", "Massachusetts", "Boston"),
    ("East", "District of Columbia", "Washington D.C."),
    ("South", "Florida", "Miami"),
    ("South", "Georgia", "Atlanta"),
    ("North", "Minnesota", "Minneapolis"),
    ("North", "Michigan", "Detroit")
]

sales_reps = [
    "Alexander Vance", "Marcus Chen", "Sophia Reynolds", "Elena Rostova", "Liam Gallagher",
    "Sarah Jenkins", "Devon Brooks", "Tariq Mansour", "Hannah Schmidt", "Carlos Mendoza",
    "Priya Patel", "Kenji Sato", "Rachel Green", "Oliver King", "Maya Lin",
    "David Rossi", "Grace Hopper", "Ethan Hunt", "Zoe Saldana", "Lucas Scott"
]

customer_names = [
    "Apex Solutions", "Quantum Systems", "Nexus Technologies", "Synergy Global",
    "Horizon Logistics", "BluePeak Capital", "Starlight Media", "Vanguard Cloud",
    "Zenith Biosystems", "Astra Corp", "Pulse Dynamics", "Cascade Fintech",
    "Silverline Retail", "OmniNet Partners", "Pinnacle Energy", "Aurora Interactive",
    "Atlas DataWorks", "Solaria Energy", "Beacon Labs", "Crestview Capital",
    "Borealis Health", "Evergreen Retail", "Helios Aerospace", "Ironclad AI",
    "Krypton Robotics", "Lumina Studio", "Meridian Consulting", "Novus Financial"
]

segments = ["Enterprise", "SMB", "Corporate", "Government / Edu"]
order_statuses = ["Completed", "Shipped", "Processing", "Delivered"]
payment_modes = ["Corporate Invoice", "Credit Card", "Direct Bank Wire", "Corporate ACH"]
ship_modes = ["Next Day Priority", "Standard Ground", "Express Air", "Freight LTL"]

start_date = datetime.date(2025, 1, 5)
end_date = datetime.date(2026, 12, 28)
total_days = (end_date - start_date).days

records = []
num_records = 650

for i in range(1, num_records + 1):
    order_id = f"ORD25-{i:04d}" if i <= 325 else f"ORD26-{i:04d}"
    random_days = int(random.triangular(0, total_days, total_days * 0.55))
    dt = start_date + datetime.timedelta(days=random_days)
    yr = dt.year
    qtr = f"Q{(dt.month - 1) // 3 + 1}"
    mo = dt.strftime("%B")
    mo_no = dt.month
    mo_yr = dt.strftime("%b-%y")
    
    cust_name = random.choice(customer_names)
    cust_id = f"CUST-{(hash(cust_name) % 900) + 100}"
    cust_seg = random.choice(segments)
    
    prod = random.choice(products)
    cat, subcat, brand, prod_name, base_price, cost_ratio = prod
    prod_id = f"SKU-{abs(hash(prod_name)) % 9000 + 1000}"
    
    reg, st, city = random.choice(regions_data)
    rep = random.choice(sales_reps)
    
    if cat == "Enterprise Hardware":
        qty = random.choices([1, 2, 3, 4, 5], weights=[40, 30, 15, 10, 5])[0]
    elif cat == "Accessories & Power":
        qty = random.randint(2, 12)
    else:
        qty = random.randint(1, 6)
        
    unit_price = round(base_price * random.uniform(0.96, 1.04), 2)
    disc_pct = random.choices([0, 5, 8, 10, 12, 15, 18, 20], weights=[25, 20, 18, 15, 10, 6, 4, 2])[0]
    
    cost_price = round(unit_price * cost_ratio * random.uniform(0.95, 1.05), 2)
    
    sales = round(qty * unit_price * (1.0 - disc_pct / 100.0), 2)
    cost = round(qty * cost_price, 2)
    profit = round(sales - cost, 2)
    margin = round((profit / sales) * 100.0, 2) if sales > 0 else 0.0
    
    status = random.choices(order_statuses, weights=[45, 35, 15, 5])[0]
    pay = random.choice(payment_modes)
    ship = random.choice(ship_modes)
    
    is_ret = random.random() < 0.135
    ret_status = "Yes" if is_ret else "No"
    
    records.append({
        "Order ID": order_id,
        "Order Date": dt.strftime("%Y-%m-%d"),
        "Year": yr,
        "Quarter": qtr,
        "Month": mo,
        "Month No": mo_no,
        "Month-Year": mo_yr,
        "Customer ID": cust_id,
        "Customer Name": cust_name,
        "Customer Segment": cust_seg,
        "Product ID": prod_id,
        "Product Name": prod_name,
        "Category": cat,
        "Sub Category": subcat,
        "Brand": brand,
        "Region": reg,
        "State": st,
        "City": city,
        "Sales Person": rep,
        "Quantity": qty,
        "Unit Price": unit_price,
        "Discount %": disc_pct,
        "Cost Price": cost_price,
        "Sales": sales,
        "Profit": profit,
        "Profit Margin %": margin,
        "Gross Profit": profit,
        "Net Profit": profit,
        "Order Status": status,
        "Payment Mode": pay,
        "Ship Mode": ship,
        "Return Status": ret_status
    })

records.sort(key=lambda r: r["Order Date"])

wb = openpyxl.Workbook()

# Sheet 1: Executive Summary
ws_sum = wb.active
ws_sum.title = "Executive Summary"
ws_sum.views.sheetView[0].showGridLines = True

# Sheet 2: Data
ws_data = wb.create_sheet(title="Data")
ws_data.views.sheetView[0].showGridLines = True

headers = list(records[0].keys())
header_fill = PatternFill(start_color="0F172A", end_color="0F172A", fill_type="solid")
header_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
border_thin = Border(
    left=Side(style='thin', color='E2E8F0'),
    right=Side(style='thin', color='E2E8F0'),
    top=Side(style='thin', color='E2E8F0'),
    bottom=Side(style='thin', color='E2E8F0')
)

ws_data.append(headers)
for col_idx, col_name in enumerate(headers, 1):
    cell = ws_data.cell(row=1, column=col_idx)
    cell.fill = header_fill
    cell.font = header_font
    cell.alignment = Alignment(horizontal="center", vertical="center")

row_alt_fill = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")

for r_idx, r in enumerate(records, 2):
    row_vals = [r[h] for h in headers]
    ws_data.append(row_vals)
    for c_idx in range(1, len(headers) + 1):
        cell = ws_data.cell(row=r_idx, column=c_idx)
        cell.font = Font(name="Segoe UI", size=10)
        cell.border = border_thin
        if r_idx % 2 == 1:
            cell.fill = row_alt_fill
            
        header = headers[c_idx - 1]
        if header in ["Unit Price", "Cost Price", "Sales", "Profit", "Gross Profit", "Net Profit"]:
            cell.number_format = "$#,##0.00"
            cell.alignment = Alignment(horizontal="right")
        elif header in ["Discount %", "Profit Margin %"]:
            cell.number_format = "0.0"
            cell.alignment = Alignment(horizontal="right")
        elif header in ["Quantity", "Month No", "Year"]:
            cell.alignment = Alignment(horizontal="center")
        elif header in ["Order ID", "Order Date", "Quarter", "Return Status", "Order Status"]:
            cell.alignment = Alignment(horizontal="center")

for col in ws_data.columns:
    max_len = max(len(str(cell.value or '')) for cell in col)
    col_letter = get_column_letter(col[0].column)
    ws_data.column_dimensions[col_letter].width = max(max_len + 3, 12)

ws_data.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{len(records) + 1}"

# Summary sheet
title_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
title_font = Font(name="Segoe UI", size=16, bold=True, color="FFFFFF")
subtitle_font = Font(name="Segoe UI", size=10, italic=True, color="64748B")

ws_sum.merge_cells("B2:G2")
t_cell = ws_sum.cell(row=2, column=2, value="EXECUTIVE SALES & PROFIT REPORT | FY 2025-2026")
t_cell.fill = title_fill
t_cell.font = title_font
t_cell.alignment = Alignment(horizontal="center", vertical="center")
ws_sum.row_dimensions[2].height = 35

ws_sum.cell(row=3, column=2, value="Generated from Enterprise Omnichannel Sales Dataset (650 Transactions)")
ws_sum.cell(row=3, column=2).font = subtitle_font

tot_sales = sum(r["Sales"] for r in records)
tot_profit = sum(r["Profit"] for r in records)
tot_margin = (tot_profit / tot_sales * 100.0) if tot_sales else 0.0
tot_orders = len(records)
tot_units = sum(r["Quantity"] for r in records)
aov = tot_sales / tot_orders
returns = sum(1 for r in records if r["Return Status"] == "Yes")
ret_rate = returns / tot_orders * 100.0

kpis = [
    ("Total Revenue (Sales)", tot_sales, "$#,##0.00"),
    ("Total Net Profit", tot_profit, "$#,##0.00"),
    ("Overall Profit Margin %", tot_margin, "0.0%"),
    ("Total Transactions", tot_orders, "#,##0"),
    ("Total Units Sold", tot_units, "#,##0"),
    ("Average Order Value (AOV)", aov, "$#,##0.00"),
    ("Return Rate %", ret_rate, "0.0%")
]

ws_sum.cell(row=5, column=2, value="Key Performance Indicator").font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
ws_sum.cell(row=5, column=3, value="Value").font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
ws_sum.cell(row=5, column=2).fill = PatternFill(start_color="0F172A", end_color="0F172A", fill_type="solid")
ws_sum.cell(row=5, column=3).fill = PatternFill(start_color="0F172A", end_color="0F172A", fill_type="solid")

for idx, (label, val, num_fmt) in enumerate(kpis, 6):
    c1 = ws_sum.cell(row=idx, column=2, value=label)
    c2 = ws_sum.cell(row=idx, column=3, value=val if "%" not in num_fmt else val / 100.0)
    c1.font = Font(name="Segoe UI", size=10, bold=(idx in [6, 7]))
    c2.font = Font(name="Segoe UI", size=10, bold=True)
    c1.border = border_thin
    c2.border = border_thin
    c2.number_format = num_fmt
    c2.alignment = Alignment(horizontal="right")
    if idx % 2 == 1:
        c1.fill = row_alt_fill
        c2.fill = row_alt_fill

ws_sum.cell(row=15, column=2, value="Category").font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
ws_sum.cell(row=15, column=3, value="Revenue ($)").font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
ws_sum.cell(row=15, column=4, value="Profit ($)").font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
ws_sum.cell(row=15, column=5, value="Margin %").font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
for c in range(2, 6):
    ws_sum.cell(row=15, column=c).fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")

cat_map = {}
for r in records:
    c = r["Category"]
    if c not in cat_map: cat_map[c] = {"sales": 0, "profit": 0}
    cat_map[c]["sales"] += r["Sales"]
    cat_map[c]["profit"] += r["Profit"]

sorted_cats = sorted(cat_map.items(), key=lambda x: x[1]["sales"], reverse=True)
for idx, (cat_name, vals) in enumerate(sorted_cats, 16):
    m = vals["profit"] / vals["sales"] if vals["sales"] else 0
    c1 = ws_sum.cell(row=idx, column=2, value=cat_name)
    c2 = ws_sum.cell(row=idx, column=3, value=vals["sales"])
    c3 = ws_sum.cell(row=idx, column=4, value=vals["profit"])
    c4 = ws_sum.cell(row=idx, column=5, value=m)
    for c in [c1, c2, c3, c4]:
        c.font = Font(name="Segoe UI", size=10)
        c.border = border_thin
    c2.number_format = "$#,##0.00"
    c3.number_format = "$#,##0.00"
    c4.number_format = "0.0%"

ws_sum.column_dimensions["B"].width = 30
ws_sum.column_dimensions["C"].width = 20
ws_sum.column_dimensions["D"].width = 20
ws_sum.column_dimensions["E"].width = 16
ws_sum.column_dimensions["F"].width = 16
ws_sum.column_dimensions["G"].width = 16

filename = "New_Sales_Profit_Dataset_2025_2026.xlsx"
for d in dirs:
    os.makedirs(d, exist_ok=True)
    out_file = os.path.join(d, filename)
    wb.save(out_file)
    print(f"Saved: {out_file} (Size: {os.path.getsize(out_file)} bytes)")

print(f"SUCCESS: Generated {len(records)} records. Total Sales: ${tot_sales:,.2f}, Total Profit: ${tot_profit:,.2f}, Margin: {tot_margin:.1f}%")
